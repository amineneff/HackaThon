# backend/document_processor.py

from pathlib import Path
from typing import Dict, List, Optional
import PyPDF2
import pdfplumber
from PIL import Image
import pytesseract
import whisper
from pdf2image import convert_from_path
import docx
from pptx import Presentation
import fitz  # PyMuPDF for better PDF image extraction
import moviepy.editor as mp
import shutil

class MultiModalDocumentProcessor:
    """
    Processes all document types including embedded media
    Supports: PDF, DOCX, PPTX, TXT, MD, Images, Audio, Video
    """
    
    def __init__(self, extract_media_dir: str = "extracted_media"):
        self.extract_dir = Path(extract_media_dir)
        self.extract_dir.mkdir(exist_ok=True)
        (self.extract_dir / "images").mkdir(exist_ok=True)
        (self.extract_dir / "audio").mkdir(exist_ok=True)
        self._whisper = None
    
    def process(self, file_path: str) -> Dict:
        """
        Main entry point - returns structured multi-modal data
        
        Returns:
            {
                'text': str,
                'images': List[str],  # paths to extracted images
                'audio': List[str],   # paths to extracted audio
                'metadata': dict
            }
        """
        ext = Path(file_path).suffix.lower()
        filename = Path(file_path).stem
        
        result = {
            'text': '',
            'images': [],
            'audio': [],
            'metadata': {
                'source_file': Path(file_path).name,
                'file_type': ext
            }
        }
        
        # Route to appropriate processor
        if ext == '.txt' or ext == '.md':
            result['text'] = self._process_text(file_path)
            
        elif ext == '.pdf':
            result.update(self._process_pdf(file_path, filename))
            
        elif ext == '.docx':
            result.update(self._process_docx(file_path, filename))
            
        elif ext == '.pptx':
            result.update(self._process_pptx(file_path, filename))
            
        elif ext in ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.webp']:
            result.update(self._process_image(file_path, filename))
            
        elif ext in ['.mp3', '.wav', '.m4a', '.flac', '.ogg']:
            result.update(self._process_audio(file_path, filename))
            
        elif ext in ['.mp4', '.avi', '.mov', '.mkv']:
            result.update(self._process_video(file_path, filename))
            
        else:
            raise ValueError(f"Unsupported file type: {ext}")
        
        return result
    
    # ============= TEXT FILES =============
    
    def _process_text(self, path: str) -> str:
        """Process plain text/markdown files"""
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()
    
    # ============= PDF FILES =============
    
    def _process_pdf(self, path: str, filename: str) -> Dict:
        """Extract text and images from PDF"""
        text_parts = []
        images = []
        
        # Extract text with pdfplumber
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        
        # Extract images with PyMuPDF (better image extraction)
        pdf_document = fitz.open(path)
        for page_num, page in enumerate(pdf_document):
            image_list = page.get_images()
            
            for img_index, img in enumerate(image_list):
                xref = img[0]
                base_image = pdf_document.extract_image(xref)
                image_bytes = base_image["image"]
                
                # Save extracted image
                img_filename = f"{filename}_page{page_num+1}_img{img_index+1}.png"
                img_path = self.extract_dir / "images" / img_filename
                
                with open(img_path, "wb") as img_file:
                    img_file.write(image_bytes)
                
                images.append(str(img_path))
        
        return {
            'text': '\n\n'.join(text_parts),
            'images': images
        }
    
    # ============= WORD DOCUMENTS =============
    
    def _process_docx(self, path: str, filename: str) -> Dict:
        """Extract text and images from Word documents"""
        doc = docx.Document(path)
        text_parts = []
        images = []
        
        # Extract text
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Extract images
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                img_index = len(images)
                img_filename = f"{filename}_img{img_index+1}.png"
                img_path = self.extract_dir / "images" / img_filename
                
                with open(img_path, 'wb') as img_file:
                    img_file.write(rel.target_part.blob)
                
                images.append(str(img_path))
        
        return {
            'text': '\n\n'.join(text_parts),
            'images': images
        }
    
    # ============= POWERPOINT =============
    
    def _process_pptx(self, path: str, filename: str) -> Dict:
        """Extract text and images from PowerPoint"""
        prs = Presentation(path)
        text_parts = []
        images = []
        
        for slide_num, slide in enumerate(prs.slides):
            # Extract text
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_parts.append(shape.text)
                
                # Extract images
                if shape.shape_type == 13:  # Picture
                    image = shape.image
                    img_filename = f"{filename}_slide{slide_num+1}_img{len(images)+1}.png"
                    img_path = self.extract_dir / "images" / img_filename
                    
                    with open(img_path, 'wb') as img_file:
                        img_file.write(image.blob)
                    
                    images.append(str(img_path))
        
        return {
            'text': '\n\n'.join(text_parts),
            'images': images
        }
    
    # ============= IMAGES =============
    
    def _process_image(self, path: str, filename: str) -> Dict:
        """Process standalone images with OCR"""
        # Copy to extracted_media for consistency
        img_path = self.extract_dir / "images" / f"{filename}{Path(path).suffix}"
        shutil.copy(path, img_path)
        
        # OCR the image
        image = Image.open(path)
        text = pytesseract.image_to_string(image)
        
        return {
            'text': text,
            'images': [str(img_path)]
        }
    
    # ============= AUDIO =============
    
    def _process_audio(self, path: str, filename: str) -> Dict:
        """Transcribe audio files"""
        # Copy to extracted_media
        audio_path = self.extract_dir / "audio" / f"{filename}{Path(path).suffix}"
        shutil.copy(path, audio_path)
        
        # Transcribe
        if self._whisper is None:
            self._whisper = whisper.load_model("base")
        
        result = self._whisper.transcribe(path)
        
        return {
            'text': result["text"],
            'audio': [str(audio_path)],
            'metadata': {
                'language': result.get('language', 'unknown')
            }
        }
    
    # ============= VIDEO =============
    
    def _process_video(self, path: str, filename: str) -> Dict:
        """Extract audio from video and transcribe"""
        # Extract audio from video
        video = mp.VideoFileClip(path)
        audio_filename = f"{filename}_audio.mp3"
        audio_path = self.extract_dir / "audio" / audio_filename
        
        video.audio.write_audiofile(str(audio_path), logger=None)
        video.close()
        
        # Transcribe extracted audio
        if self._whisper is None:
            self._whisper = whisper.load_model("base")
        
        result = self._whisper.transcribe(str(audio_path))
        
        return {
            'text': result["text"],
            'audio': [str(audio_path)],
            'metadata': {
                'language': result.get('language', 'unknown'),
                'source_type': 'video_audio_extraction'
            }
        }